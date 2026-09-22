#!/usr/bin/env python3
"""Outreach tracker: create it, append rows to it, and stage rows for Google Sheets.

Using this script instead of hand written openpyxl keeps formatting consistent
across runs and makes the append path testable.

Create the tracker:
    python3 tracker.py create /path/to/outreach_tracker.xlsx

Append a row:
    python3 tracker.py append /path/to/outreach_tracker.xlsx \
        --firm "Example Partners" --type "Boutique IB" \
        --name "Jane Doe" --title "Managing Director" \
        --affiliation "Not verified" --url "example.com/team/jane-doe" \
        --email "jane.doe@example.com" --notes "Gmail draft created. ..." \
        --status drafted

Stage a row for a Google Sheet when no Sheets write path is available:
    python3 tracker.py pending /path/to/tracker_pending.csv --firm ... [same flags]

Dates default to today. Append refuses to add a firm that already has a row
unless --allow-duplicate is passed, which is the main guard against a rerun
double logging the same contact.
"""

import argparse
import csv
import datetime
import os
import sys

COLUMNS = [
    ("Firm", 24),
    ("Firm Type", 18),
    ("Name", 20),
    ("Title", 28),
    ("Affiliation", 34),
    ("Profile URL", 28),
    ("Email", 28),
    ("Notes", 90),
    ("Date Drafted", 14),
    ("Status", 12),
]

HEADERS = [name for name, _ in COLUMNS]


def _styles():
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

    thin = Side(style="thin", color="CCCCCC")
    return {
        "header_font": Font(name="Arial", size=10, bold=True, color="FFFFFF"),
        "header_fill": PatternFill("solid", fgColor="1F3864"),
        "body_font": Font(name="Arial", size=10),
        "row_fill": PatternFill("solid", fgColor="EEF2FF"),
        "blank_fill": PatternFill(),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "align": Alignment(vertical="center", wrap_text=True),
        "header_align": Alignment(vertical="center", horizontal="center", wrap_text=True),
    }


def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        sys.exit(
            "openpyxl is required. Install it with:\n"
            "    pip install openpyxl --break-system-packages"
        )


def cmd_create(args):
    _require_openpyxl()
    import openpyxl

    if os.path.exists(args.path):
        sys.exit(f"Refusing to overwrite existing file: {args.path}")

    s = _styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Outreach"

    for idx, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = s["header_font"]
        cell.fill = s["header_fill"]
        cell.border = s["border"]
        cell.alignment = s["header_align"]
        ws.column_dimensions[cell.column_letter].width = width

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    wb.save(args.path)

    print(f"Created tracker: {args.path}")
    print("Columns: " + ", ".join(HEADERS))


def _row_values(args):
    return [
        args.firm,
        args.type,
        args.name,
        args.title,
        args.affiliation,
        args.url,
        args.email,
        args.notes,
        args.date or datetime.date.today().isoformat(),
        args.status,
    ]


def cmd_append(args):
    _require_openpyxl()
    import openpyxl

    if not os.path.exists(args.path):
        sys.exit(f"Tracker not found: {args.path}\nRun: tracker.py create {args.path}")

    s = _styles()
    wb = openpyxl.load_workbook(args.path)
    ws = wb.active

    header = [c.value for c in ws[1]]
    if header[: len(HEADERS)] != HEADERS:
        print(f"WARNING: header mismatch.\n  expected: {HEADERS}\n  found:    {header}")
        print("Appending anyway in column order. Check the sheet afterwards.")

    if not args.allow_duplicate:
        existing = {
            str(row[0]).strip().lower()
            for row in ws.iter_rows(min_row=2, values_only=True)
            if row and row[0]
        }
        if args.firm.strip().lower() in existing:
            sys.exit(
                f"'{args.firm}' already has a row. This usually means the run is "
                "repeating work already done. Pass --allow-duplicate to force it."
            )

    new_row = ws.max_row + 1
    for col, value in enumerate(_row_values(args), start=1):
        cell = ws.cell(row=new_row, column=col, value=value)
        cell.font = s["body_font"]
        cell.fill = s["row_fill"] if new_row % 2 == 0 else s["blank_fill"]
        cell.border = s["border"]
        cell.alignment = s["align"]

    ws.row_dimensions[new_row].height = 40
    wb.save(args.path)

    print(f"Appended row {new_row}: {args.firm} / {args.name} / {args.email}")


def cmd_pending(args):
    """Stage a row as CSV when the Google Sheet could not be written.

    Nothing is lost when a Sheets write fails; the row lands here and the user
    imports it. Appends, never overwrites.
    """
    exists = os.path.exists(args.path)
    with open(args.path, "a", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        if not exists:
            writer.writerow(HEADERS)
        writer.writerow(_row_values(args))

    print(f"Staged row in {args.path}")
    print("This row is NOT in the Google Sheet yet. Import it or paste it in.")


def build_parser():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="command", required=True)

    create = sub.add_parser("create", help="create an empty formatted tracker")
    create.add_argument("path")
    create.set_defaults(func=cmd_create)

    def add_row_args(parser):
        parser.add_argument("path")
        parser.add_argument("--firm", required=True)
        parser.add_argument("--type", required=True, help="firm type, e.g. Boutique IB")
        parser.add_argument("--name", required=True)
        parser.add_argument("--title", required=True)
        parser.add_argument("--affiliation", default="Not verified")
        parser.add_argument("--url", default="")
        parser.add_argument("--email", required=True)
        parser.add_argument("--notes", default="")
        parser.add_argument("--date", default="", help="YYYY-MM-DD, defaults to today")
        parser.add_argument("--status", default="drafted", choices=["drafted", "sent", "replied", "closed"])

    append = sub.add_parser("append", help="append one contact row")
    add_row_args(append)
    append.add_argument("--allow-duplicate", action="store_true", help="permit a firm that already has a row")
    append.set_defaults(func=cmd_append)

    pending = sub.add_parser("pending", help="stage a row as CSV for a Google Sheet")
    add_row_args(pending)
    pending.set_defaults(func=cmd_pending)

    return p


if __name__ == "__main__":
    args = build_parser().parse_args()
    args.func(args)
